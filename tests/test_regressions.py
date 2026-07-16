import os
import tempfile
import unittest

import app as egg


def _all_report_items(quantity_overrides=None):
    quantity_overrides = quantity_overrides or {}
    items = []
    for category, specs in egg.PRESET_TEMPLATES.items():
        for spec in specs:
            key = f"{category}_{spec}"
            items.append({
                "category": category,
                "spec": spec,
                "quantity": quantity_overrides.get(key, 0),
            })
    return items


class EggNumRegressionTests(unittest.TestCase):
    def setUp(self):
        fd, self.db_path = tempfile.mkstemp(prefix="eggnum_test_", suffix=".db")
        os.close(fd)
        os.remove(self.db_path)

        egg.app.config["DATABASE"] = self.db_path
        egg.app.config["TESTING"] = True
        with egg.app.app_context():
            egg.init_db()
        self.client = egg.app.test_client()

    def tearDown(self):
        try:
            os.remove(self.db_path)
        except FileNotFoundError:
            pass

    def test_stale_record_id_cannot_merge_into_another_date(self):
        first_category = next(iter(egg.PRESET_TEMPLATES))
        first_spec = egg.PRESET_TEMPLATES[first_category][0]
        first_key = f"{first_category}_{first_spec}"
        items = _all_report_items({first_key: 5})

        created = self.client.post("/api/submit", json={
            "store_name": egg.DEFAULT_STORE_NAME,
            "record_date": "2026-06-20",
            "items": items,
        })
        self.assertEqual(created.status_code, 200)
        record_id = created.json["record_id"]

        stale_merge = self.client.post("/api/submit", json={
            "store_name": egg.DEFAULT_STORE_NAME,
            "record_date": "2026-06-21",
            "record_id": record_id,
            "merge": True,
            "items": [{"category": first_category, "spec": first_spec, "quantity": 9}],
        })
        self.assertEqual(stale_merge.status_code, 409)

        full_save = self.client.post("/api/submit", json={
            "store_name": egg.DEFAULT_STORE_NAME,
            "record_date": "2026-06-21",
            "record_id": record_id,
            "items": _all_report_items({first_key: 9}),
        })
        self.assertEqual(full_save.status_code, 200)

        history = self.client.get("/api/history").json
        self.assertEqual(
            sorted(record["record_date"] for record in history),
            ["2026-06-20", "2026-06-21"],
        )

    def test_attendance_history_record_can_be_updated(self):
        created = self.client.post("/api/attendance", json={
            "record_date": "2026-06-29",
            "time_start": "09:00",
            "time_end": "11:00",
            "hours": 2,
            "note": "old",
        })
        self.assertEqual(created.status_code, 200)

        updated = self.client.put(f"/api/attendance/{created.json['id']}", json={
            "record_date": "2026-06-28",
            "time_start": "10:30",
            "time_end": "13:00",
            "hours": 2.5,
            "note": "fixed",
        })
        self.assertEqual(updated.status_code, 200)

        groups = self.client.get("/api/attendance-history").json["groups"]
        self.assertEqual(len(groups), 1)
        self.assertEqual(groups[0]["date"], "2026-06-28")
        self.assertEqual(groups[0]["total"], 2.5)
        self.assertEqual(groups[0]["entries"][0]["note"], "fixed")

    def test_attendance_export_resets_single_segment_row_height(self):
        import io
        import openpyxl

        first = self.client.post("/api/attendance", json={
            "record_date": "2026-06-29",
            "time_start": "09:00",
            "time_end": "11:00",
            "hours": 2,
            "note": "",
        })
        self.assertEqual(first.status_code, 200)

        second = self.client.post("/api/attendance", json={
            "record_date": "2026-06-29",
            "time_start": "13:00",
            "time_end": "15:00",
            "hours": 2,
            "note": "",
        })
        self.assertEqual(second.status_code, 200)

        deleted = self.client.delete(f"/api/attendance/{second.json['id']}")
        self.assertEqual(deleted.status_code, 200)

        exported = self.client.get("/api/attendance/export?from=2026-06-29&to=2026-06-29")
        self.assertEqual(exported.status_code, 200)

        wb = openpyxl.load_workbook(io.BytesIO(exported.data))
        ws = wb.active
        self.assertEqual(ws.row_dimensions[2].height, 15)

    def test_attendance_equal_start_end_is_zero_hours(self):
        created = self.client.post("/api/attendance", json={
            "record_date": "2026-06-29",
            "time_start": "00:00",
            "time_end": "00:00",
            "hours": 24,
            "note": "休息",
        })
        self.assertEqual(created.status_code, 200)

        groups = self.client.get("/api/attendance-history").json["groups"]
        self.assertEqual(groups[0]["total"], 0)
        self.assertEqual(groups[0]["entries"][0]["hours"], 0)
        self.assertEqual(groups[0]["entries"][0]["note"], "休息")

        import io
        import openpyxl

        exported = self.client.get("/api/attendance/export?from=2026-06-29&to=2026-06-29")
        self.assertEqual(exported.status_code, 200)

        wb = openpyxl.load_workbook(io.BytesIO(exported.data))
        ws = wb.active
        self.assertEqual(ws.cell(row=2, column=2).value, "-")
        self.assertEqual(ws.cell(row=2, column=3).value, 0)

    def test_reserve_cannot_go_below_zero(self):
        response = self.client.post("/api/reserve", json={
            "category": "农家蛋",
            "spec": 30,
            "delta": -1,
        })
        self.assertEqual(response.status_code, 400)
        self.assertEqual(response.json["success"], False)

    def test_today_endpoint_uses_latest_record_shape(self):
        first_category = next(iter(egg.PRESET_TEMPLATES))
        first_spec = egg.PRESET_TEMPLATES[first_category][0]
        first_key = f"{first_category}_{first_spec}"
        items = _all_report_items({first_key: 3})

        created = self.client.post("/api/submit", json={
            "store_name": egg.DEFAULT_STORE_NAME,
            "record_date": "2026-06-30",
            "items": items,
        })
        self.assertEqual(created.status_code, 200)

        today = self.client.get("/api/today?date=2026-06-30")
        self.assertEqual(today.status_code, 200)
        self.assertTrue(today.json["found"])
        self.assertEqual(today.json["record_id"], created.json["record_id"])
        self.assertEqual(today.json["items"][0]["quantity"], 3)

    def test_reserve_api_returns_items_array(self):
        response = self.client.get("/api/reserve")
        self.assertEqual(response.status_code, 200)
        self.assertIn("items", response.json)
        self.assertIsInstance(response.json["items"], list)


if __name__ == "__main__":
    unittest.main()
